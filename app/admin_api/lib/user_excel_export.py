from __future__ import annotations

from datetime import UTC, date, datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FILL = PatternFill("solid", fgColor="1F4A3A")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=16, color="1F4A3A")
META_FONT = Font(size=11, color="444444")
THIN = Side(style="thin", color="D4D4D4")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
ALT_FILL = PatternFill("solid", fgColor="F7FAF8")


def _fmt_dt(value: datetime | None) -> str:
    if not value:
        return ""
    if value.tzinfo is None:
        value = value.replace(tzinfo=UTC)
    return value.astimezone(UTC).strftime("%d %b %Y, %H:%M UTC")


def _fmt_date(value: date | None) -> str:
    return value.strftime("%d %b %Y") if value else ""


def _fmt_paise(paise: int | None) -> str:
    if not paise:
        return "0.00"
    return f"{paise / 100:,.2f}"


def _yes_no(value: bool | None) -> str:
    return "Yes" if value else "No"


def _style_header_row(ws: Worksheet, *, row: int, col_count: int) -> None:
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = CELL_BORDER
    ws.row_dimensions[row].height = 28


def _autosize_columns(ws: Worksheet, *, min_width: int = 10, max_width: int = 42) -> None:
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = 0
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value is None:
                    continue
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max(max_len + 2, min_width), max_width)


def build_user_workbook(
    *,
    customers: list[dict[str, Any]],
    addresses: list[dict[str, Any]],
    stats: dict[str, int],
    filters: dict[str, str | None],
    exported_at: datetime,
) -> bytes:
    wb = Workbook()
    overview = wb.active
    overview.title = "Overview"

    overview["A1"] = "Chic A Boo — Customer export"
    overview["A1"].font = TITLE_FONT
    overview["A2"] = f"Generated {_fmt_dt(exported_at)}"
    overview["A2"].font = META_FONT

    filter_bits = []
    if filters.get("search"):
        filter_bits.append(f'Search: "{filters["search"]}"')
    if filters.get("status"):
        filter_bits.append(f'Status: {filters["status"]}')
    overview["A3"] = "Filters: " + (", ".join(filter_bits) if filter_bits else "All customers")
    overview["A3"].font = META_FONT

    summary_rows = [
        ("Total customers", stats.get("total", 0)),
        ("Active", stats.get("active", 0)),
        ("Suspended", stats.get("suspended", 0)),
        ("Blocked", stats.get("blocked", 0)),
        ("Pending verification", stats.get("pending_verification", 0)),
        ("New in last 30 days", stats.get("new_30d", 0)),
        ("With at least one order", stats.get("with_orders", 0)),
        ("Saved addresses", len(addresses)),
    ]
    start = 5
    overview.cell(row=start, column=1, value="Metric").font = Font(bold=True)
    overview.cell(row=start, column=2, value="Count").font = Font(bold=True)
    for offset, (label, count) in enumerate(summary_rows, start=1):
        overview.cell(row=start + offset, column=1, value=label)
        overview.cell(row=start + offset, column=2, value=count)

    overview.column_dimensions["A"].width = 28
    overview.column_dimensions["B"].width = 14

    customer_headers = [
        "Customer #",
        "Full name",
        "Email",
        "Phone",
        "Status",
        "Status reason",
        "Email verified",
        "Phone verified",
        "Gender",
        "Date of birth",
        "Loyalty points",
        "Orders",
        "Lifetime spend (₹)",
        "Last order",
        "Last login",
        "Joined",
        "Default city",
        "Default state",
        "Default pincode",
        "Email marketing",
        "SMS marketing",
        "Order updates (email)",
        "Order updates (SMS)",
        "Push notifications",
        "Language",
        "Currency",
    ]
    ws = wb.create_sheet("Customers")
    ws.append(customer_headers)
    _style_header_row(ws, row=1, col_count=len(customer_headers))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(customer_headers))}1"

    for idx, row in enumerate(customers, start=2):
        ws.append(
            [
                row.get("customer_number"),
                row.get("full_name") or "",
                row.get("email") or "",
                row.get("phone") or "",
                row.get("status") or "",
                row.get("status_reason") or "",
                _yes_no(row.get("email_verified")),
                _yes_no(row.get("phone_verified")),
                row.get("gender") or "",
                _fmt_date(row.get("date_of_birth")),
                row.get("loyalty_points") or 0,
                row.get("order_count") or 0,
                _fmt_paise(row.get("total_spent_paise")),
                _fmt_dt(row.get("last_order_at")),
                _fmt_dt(row.get("last_login_at")),
                _fmt_dt(row.get("created_at")),
                row.get("default_city") or "",
                row.get("default_state") or "",
                row.get("default_postal_code") or "",
                _yes_no(row.get("email_marketing")),
                _yes_no(row.get("sms_marketing")),
                _yes_no(row.get("order_updates_email")),
                _yes_no(row.get("order_updates_sms")),
                _yes_no(row.get("push_notifications")),
                row.get("preferred_language") or "",
                row.get("currency") or "",
            ]
        )
        if idx % 2 == 0:
            for col in range(1, len(customer_headers) + 1):
                ws.cell(row=idx, column=col).fill = ALT_FILL
        for col in range(1, len(customer_headers) + 1):
            ws.cell(row=idx, column=col).border = CELL_BORDER

    _autosize_columns(ws)

    address_headers = [
        "Customer #",
        "Email",
        "Label",
        "Recipient",
        "Phone",
        "Line 1",
        "Line 2",
        "Landmark",
        "City",
        "State",
        "Postal code",
        "Country",
        "Default",
        "Type",
    ]
    addr_ws = wb.create_sheet("Addresses")
    addr_ws.append(address_headers)
    _style_header_row(addr_ws, row=1, col_count=len(address_headers))
    addr_ws.freeze_panes = "A2"
    addr_ws.auto_filter.ref = f"A1:{get_column_letter(len(address_headers))}1"

    for idx, row in enumerate(addresses, start=2):
        addr_ws.append(
            [
                row.get("customer_number"),
                row.get("email") or "",
                row.get("label") or "",
                row.get("full_name") or "",
                row.get("phone") or "",
                row.get("line1") or "",
                row.get("line2") or "",
                row.get("landmark") or "",
                row.get("city") or "",
                row.get("state") or "",
                row.get("postal_code") or "",
                row.get("country") or "",
                _yes_no(row.get("is_default")),
                row.get("address_type") or "",
            ]
        )
        if idx % 2 == 0:
            for col in range(1, len(address_headers) + 1):
                addr_ws.cell(row=idx, column=col).fill = ALT_FILL
        for col in range(1, len(address_headers) + 1):
            addr_ws.cell(row=idx, column=col).border = CELL_BORDER

    _autosize_columns(addr_ws)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
